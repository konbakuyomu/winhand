"""Turn files into something a model can look at: images (resized to what vision models
use), PDF text or rendered pages, and the text of Word, PowerPoint and Excel files."""

from __future__ import annotations

import io
import os
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

MAX_SIDE = 1568  # longest edge vision models work with; larger images are only scaled down
# Encoded size budget for one image in a tool result. Claude Desktop rejects tool results over
# 1 MB, and the image travels base64-encoded (x4/3) inside JSON, so stay well below: 600 KB
# of image is about 800 KB on the wire. PNG keeps screen text sharp; when a busy screen does
# not fit, JPEG quality steps down, then the image is scaled down.
IMAGE_BUDGET = 600_000

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".ico"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".docm", ".pptx", ".pptm", ".xlsx", ".xlsm"}
LEGACY_OFFICE = {".doc", ".xls", ".ppt"}


class MediaError(RuntimeError):
    pass


def kind_of(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in DOCUMENT_EXTENSIONS:
        return "document"
    if ext in LEGACY_OFFICE:
        return "legacy_office"
    return "other"


# ------------------------------------------------------------------ images


def _fit_budget(image, budget: int):
    """Smallest acceptable encoding: PNG if it fits, else JPEG at falling quality, else smaller."""
    from PIL import Image

    alpha = image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info)
    while True:
        buffer = io.BytesIO()
        if alpha:
            image.convert("RGBA").save(buffer, "PNG", optimize=True)
            if buffer.tell() <= budget or max(image.size) <= 256:
                return image, buffer, "png"
        else:
            rgb = image.convert("RGB")
            rgb.save(buffer, "PNG", optimize=True)
            if buffer.tell() <= budget:
                return image, buffer, "png"
            for quality in (85, 75, 60):
                buffer = io.BytesIO()
                rgb.save(buffer, "JPEG", quality=quality, optimize=True)
                if buffer.tell() <= budget or max(image.size) <= 256:
                    return image, buffer, "jpeg"
        image = image.resize(
            (max(1, round(image.size[0] * 0.8)), max(1, round(image.size[1] * 0.8))), Image.Resampling.LANCZOS
        )


def encode_image(image, max_side: int = MAX_SIDE, budget: int = IMAGE_BUDGET) -> tuple[bytes, str, dict]:
    """(bytes, format, facts) for a PIL image: at most `max_side` on the longest edge and at
    most `budget` bytes encoded (see IMAGE_BUDGET)."""
    from PIL import Image, ImageOps

    if getattr(image, "n_frames", 1) > 1:
        image.seek(0)
    image = ImageOps.exif_transpose(image)
    original = image.size
    scale = min(1.0, max_side / max(original))
    if scale < 1.0:
        image = image.resize(
            (max(1, round(original[0] * scale)), max(1, round(original[1] * scale))), Image.Resampling.LANCZOS
        )
    image, buffer, fmt = _fit_budget(image, budget)
    facts = {
        "width": image.size[0],
        "height": image.size[1],
        "original_width": original[0],
        "original_height": original[1],
        "scale": round(image.size[0] / original[0], 4),
        "format": fmt,
        "bytes": buffer.tell(),
    }
    return buffer.getvalue(), fmt, facts


def read_image(
    path: str, max_side: int = MAX_SIDE, region: list[int] | None = None
) -> tuple[bytes, str, dict]:
    """An image file, optionally only `region` = [left, top, width, height] of it (in the file's
    own pixels): a zoom into a full-resolution screenshot shows small text sharply."""
    from PIL import Image, UnidentifiedImageError

    try:
        with Image.open(path) as image:
            if region is not None:
                if len(region) != 4 or region[2] <= 0 or region[3] <= 0:
                    raise MediaError("region is [left, top, width, height] in the image's pixels")
                left, top, width, height = region
                box = (
                    max(0, left),
                    max(0, top),
                    min(image.width, left + width),
                    min(image.height, top + height),
                )
                if box[0] >= box[2] or box[1] >= box[3]:
                    raise MediaError(f"region is outside the {image.width}x{image.height} image")
                image = image.crop(box)
            data, fmt, facts = encode_image(image, max_side)
    except (UnidentifiedImageError, OSError) as exc:
        raise MediaError(f"cannot open image {path}: {exc}") from exc
    facts["path"] = path
    facts["file_bytes"] = os.path.getsize(path)
    if region is not None:
        facts["region"] = region
    return data, fmt, facts


# --------------------------------------------------------------- documents


def parse_pages(spec: str | None, count: int) -> list[int]:
    """'1-3,5' -> [0, 1, 2, 4] (clamped to the document)."""
    if not spec:
        return list(range(count))
    chosen: list[int] = []
    for part in spec.replace("，", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, _, end = part.partition("-")
            first = int(start or 1)
            last = int(end or count)
            chosen.extend(range(first - 1, min(last, count)))
        else:
            chosen.append(int(part) - 1)
    return [i for i in dict.fromkeys(chosen) if 0 <= i < count]


def read_document(path: str, pages: str | None = None, render: bool = False, max_chars: int = 60_000) -> dict:
    """{'text', 'facts', 'images': [(bytes, fmt, facts)]} for PDF/Word/PowerPoint/Excel."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        result = _pdf(path, pages, render)
    elif ext in (".docx", ".docm"):
        result = {"text": _docx(path), "facts": {"type": "docx"}, "images": []}
    elif ext in (".pptx", ".pptm"):
        result = _pptx(path, pages)
    elif ext in (".xlsx", ".xlsm"):
        result = _xlsx(path)
    else:
        raise MediaError(f"unsupported document type {ext}")
    text = result["text"]
    if len(text) > max_chars:
        result["facts"]["truncated_chars"] = len(text) - max_chars
        text = text[:max_chars] + f"\n\n… [{len(text) - max_chars} more characters; ask for specific pages]"
    result["text"] = text
    result["facts"]["path"] = path
    return result


def _pdf(path: str, pages: str | None, render: bool) -> dict:
    import pypdfium2 as pdfium

    try:
        document = pdfium.PdfDocument(path)
    except pdfium.PdfiumError as exc:
        raise MediaError(f"cannot open PDF (encrypted or damaged?): {exc}") from exc
    try:
        count = len(document)
        chosen = parse_pages(pages, count)
        parts, empty = [], []
        for index in chosen:
            page = document[index]
            text = page.get_textpage().get_text_range().replace("\r\n", "\n").strip()
            if len(text) < 5:  # no real text layer: a scan
                empty.append(index)
            parts.append(f"--- 第 {index + 1} 页 / page {index + 1} ---\n{text}")
        # scanned pages have no text layer: show them as pictures instead (a few at most)
        to_render = chosen if render else empty
        images = []
        for index in to_render[:6]:
            page = document[index]
            width, height = page.get_size()
            scale = min(3.0, MAX_SIDE / max(width, height))
            bitmap = page.render(scale=scale)
            data, fmt, facts = encode_image(bitmap.to_pil())
            facts["page"] = index + 1
            images.append((data, fmt, facts))
        facts = {"type": "pdf", "pages": count, "shown_pages": [i + 1 for i in chosen]}
        if empty:
            facts["pages_without_text"] = [i + 1 for i in empty]
        if len(to_render) > 6:
            facts["rendered_only_first"] = 6
        return {"text": "\n\n".join(parts), "facts": facts, "images": images}
    finally:
        document.close()


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _xml(archive: zipfile.ZipFile, name: str) -> ElementTree.Element:
    return ElementTree.fromstring(archive.read(name))


def _docx_paragraph(paragraph: ElementTree.Element) -> str:
    out = []
    for node in paragraph.iter():
        if node.tag == _W + "t" and node.text:
            out.append(node.text)
        elif node.tag == _W + "tab":
            out.append("\t")
        elif node.tag in (_W + "br", _W + "cr"):
            out.append("\n")
    text = "".join(out)
    style = paragraph.find(f"{_W}pPr/{_W}pStyle")
    level = re.search(r"(\d)$", style.get(_W + "val", "")) if style is not None else None
    if style is not None and "heading" in style.get(_W + "val", "").lower() and level and text.strip():
        return "#" * int(level.group(1)) + " " + text
    return text


def _docx(path: str) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            body = _xml(archive, "word/document.xml").find(_W + "body")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise MediaError(f"not a valid .docx: {exc}") from exc
    lines = []
    for block in body if body is not None else []:
        if block.tag == _W + "p":
            lines.append(_docx_paragraph(block))
        elif block.tag == _W + "tbl":
            for row in block.iter(_W + "tr"):
                cells = [
                    " ".join(_docx_paragraph(p) for p in cell.iter(_W + "p")).strip()
                    for cell in row.iter(_W + "tc")
                ]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _pptx(path: str, pages: str | None) -> dict:
    try:
        with zipfile.ZipFile(path) as archive:
            names = sorted(
                (n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                key=lambda n: int(re.search(r"(\d+)", n).group(1)),
            )
            chosen = parse_pages(pages, len(names))
            parts = []
            for index in chosen:
                root = _xml(archive, names[index])
                paragraphs = [
                    "".join(t.text or "" for t in p.iter(_A + "t")).strip() for p in root.iter(_A + "p")
                ]
                parts.append(
                    f"--- 第 {index + 1} 张 / slide {index + 1} ---\n" + "\n".join(x for x in paragraphs if x)
                )
    except (zipfile.BadZipFile, KeyError) as exc:
        raise MediaError(f"not a valid .pptx: {exc}") from exc
    return {"text": "\n\n".join(parts), "facts": {"type": "pptx", "slides": len(names)}, "images": []}


def _xlsx(path: str, max_rows: int = 500, max_cols: int = 50) -> dict:
    import openpyxl

    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise MediaError(f"cannot open workbook: {exc}") from exc
    parts, sheets = [], []
    try:
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                rows.append("\t".join(cells))
            while rows and not rows[-1]:
                rows.pop()
            total = sheet.max_row or 0
            more = f"\n… ({total - max_rows} more rows)" if total > max_rows else ""
            parts.append(f"--- 工作表 / sheet: {sheet.title} ({total} rows) ---\n" + "\n".join(rows) + more)
            sheets.append(sheet.title)
    finally:
        book.close()
    return {"text": "\n\n".join(parts), "facts": {"type": "xlsx", "sheets": sheets}, "images": []}
